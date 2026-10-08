"""Tạo file mẫu để test trích xuất văn bản: .docx, .pdf, .txt + file lỗi.

Chạy: python3 tests/make_fixtures.py   (ra /tmp/fixture.*)
"""
import os
import subprocess

OUT = "/tmp"
PARA = """Water scarcity is one of the most pressing challenges of the twenty-first century.
Urbanization, population growth and climate change are the main contributing factors to water shortage.
Governments should focus solely on technology to solve water issues is a controversial claim.
Both the government and each individual have a role to play in protecting water resources.
Conservation measures include recycling, reducing, reusing and refilling domestic water.
Immediate action is needed to address the growing water shortage crisis worldwide.
"""


def main():
    from docx import Document

    doc = Document()
    doc.add_heading("Water Scarcity Reading Passage", level=1)
    for _ in range(6):
        for line in PARA.strip().split("\n"):
            doc.add_paragraph(line)
    doc.save(os.path.join(OUT, "fixture.docx"))

    with open(os.path.join(OUT, "fixture.txt"), "w", encoding="utf-8") as f:
        f.write(PARA * 8)

    html = "<html><body style='font-family:serif;font-size:14px'>" + "".join(
        f"<p>{line}</p>" for line in (PARA * 6).strip().split("\n")
    ) + "</body></html>"
    html_path = os.path.join(OUT, "fixture.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    subprocess.run([
        "/usr/bin/google-chrome", "--headless=new", "--disable-gpu", "--no-sandbox",
        f"--print-to-pdf={os.path.join(OUT, 'fixture.pdf')}", f"file://{html_path}",
    ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    with open(os.path.join(OUT, "fixture.exe"), "wb") as f:
        f.write(b"MZ fake binary")
    with open(os.path.join(OUT, "short.txt"), "w", encoding="utf-8") as f:
        f.write("ngắn")

    # kiểm chứng trích xuất thật
    import sys
    sys.path.insert(0, os.getcwd())
    import main as app_main
    for name in ("fixture.docx", "fixture.pdf", "fixture.txt"):
        path = os.path.join(OUT, name)
        with open(path, "rb") as f:
            text = app_main.extract_text(path, f.read())
        print(f"extract {name} -> {len(text)} chars | {text[:50]!r}")
    try:
        app_main.extract_text("fixture.exe", b"MZ")
        print("extract .exe -> SHOULD HAVE RAISED")
    except ValueError as e:
        print("extract .exe raises ValueError:", e)


if __name__ == "__main__":
    main()
