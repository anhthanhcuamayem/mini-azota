from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, Response
from datetime import datetime, timedelta, timezone
import asyncio
import io
import json
import os
import random
import re
import subprocess
import tempfile
import uuid

import requests

import school_store as store
import supabase_store as db
from school_api import router as school_router

app = FastAPI()
app.include_router(school_router)

FAVICON_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">
  <defs>
    <linearGradient id="g" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#ff4500"/>
      <stop offset="100%" stop-color="#ff8c00"/>
    </linearGradient>
  </defs>
  <circle cx="50" cy="50" r="48" fill="url(#g)"/>
  <text x="50" y="68" font-size="50" text-anchor="middle" fill="#ffffff" font-family="system-ui, sans-serif">🔥</text>
</svg>"""

@app.get("/favicon.ico", include_in_schema=False)
async def get_favicon():
    return Response(content=FAVICON_SVG, media_type="image/svg+xml")

@app.get("/style.css")
async def get_style():
    if os.path.exists("style.css"):
        return FileResponse("style.css", media_type="text/css")
    return HTMLResponse("", status_code=404)

EXAM_DURATION_MINUTES = 60
EXAM_QUESTION_COUNT = 40
AI_CONFIG_FILE = "ai_config.json"
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MIN_SOURCE_CHARS = 200
MAX_SOURCE_CHARS = 60000


def load_env(path=".env"):
    """Đọc file .env đơn giản (không ghi đè biến môi trường đã có)."""
    if not os.path.exists(path):
        return
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value


load_env()
AI_BASE_URL = os.environ.get("AI_BASE_URL", "").strip()
AI_API_KEY = os.environ.get("AI_API_KEY", "").strip()
AI_MODEL = os.environ.get("AI_MODEL", "llama-3.3-70b-versatile").strip()
ADMIN_KEY = os.environ.get("ADMIN_KEY", "").strip()

def load_ai_config():
    """Nạp cấu hình AI do admin lưu; bí mật chỉ tồn tại phía server."""
    global AI_BASE_URL, AI_API_KEY, AI_MODEL
    if not os.path.exists(AI_CONFIG_FILE):
        return
    try:
        with open(AI_CONFIG_FILE, "r", encoding="utf-8") as f:
            config = json.load(f)
        AI_BASE_URL = str(config.get("base_url") or AI_BASE_URL).strip()
        AI_API_KEY = str(config.get("api_key") or AI_API_KEY).strip()
        AI_MODEL = str(config.get("model") or AI_MODEL).strip()
    except (OSError, ValueError, TypeError):
        pass


load_ai_config()

def chat_completions_url(base_url):
    url = (base_url or "").strip().rstrip("/")
    return url if url.endswith("/chat/completions") else url + "/chat/completions"

def save_ai_config(base_url, api_key, model):
    with open(AI_CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump({"base_url": base_url, "api_key": api_key, "model": model}, f)
    try:
        os.chmod(AI_CONFIG_FILE, 0o600)
    except OSError:
        pass

def get_vn_time():
    vn_tz = timezone(timedelta(hours=7))
    return datetime.now(vn_tz).strftime("%H:%M:%S %d/%m/%Y")

def get_current_utc():
    return datetime.now(timezone.utc)

def format_duration(seconds):
    if seconds < 0:
        seconds = 0
    mins = seconds // 60
    secs = seconds % 60
    return f"{mins:02d}:{secs:02d}"

def save_session(session_id, name, start_time, shuffled_questions=None, quiz_id=None):
    """Lưu phiên làm bài (ẩn danh) vào Supabase."""
    store.save_exam_session(session_id, name, start_time, shuffled_questions or [],
                            None, None, quiz_id=quiz_id)


def get_session(session_id):
    return store.get_exam_session(session_id)


def delete_session(session_id):
    store.delete_exam_session(session_id)


def cleanup_old_sessions(max_age_hours=2):
    store.cleanup_old_sessions(max_age_hours)


def save_submission(entry):
    store.save_attempt(entry)


def to_vn_display(iso):
    """Đổi ISO UTC sang chuỗi giờ Việt Nam để hiển thị."""
    try:
        dt = datetime.fromisoformat(iso)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(timezone(timedelta(hours=7))).strftime("%H:%M:%S %d/%m/%Y")
    except (ValueError, TypeError):
        return iso or ""


# ---------------- Kho bộ đề (sinh bằng AI từ tài liệu) ----------------
class AIError(Exception):
    """Lỗi khi gọi AI sinh đề."""


def extract_text(filename, data):
    """Trích văn bản thuần từ .txt/.md/.docx/.pdf."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    if ext in (".docx", ".doc"):
        try:
            from docx import Document
        except ImportError:
            raise ValueError("Thiếu thư viện python-docx. Chạy: pip install python-docx")
        try:
            docx_data = data
            if ext == ".doc":
                # python-docx không đọc định dạng .doc nhị phân cũ; dùng LibreOffice có sẵn trên server.
                with tempfile.TemporaryDirectory() as temp_dir:
                    input_path = os.path.join(temp_dir, "source.doc")
                    with open(input_path, "wb") as f:
                        f.write(data)
                    try:
                        subprocess.run(
                            ["libreoffice", "-env:UserInstallation=file://" + temp_dir + "/lo-profile",
                             "--headless", "--convert-to", "docx", "--outdir", temp_dir, input_path],
                            check=True, capture_output=True, timeout=45,
                        )
                    except (OSError, subprocess.SubprocessError) as e:
                        raise ValueError("Không chuyển đổi được file .doc. Hãy lưu file thành .docx rồi tải lại.") from e
                    converted_path = os.path.join(temp_dir, "source.docx")
                    if not os.path.exists(converted_path):
                        raise ValueError("Không chuyển đổi được file .doc. Hãy lưu file thành .docx rồi tải lại.")
                    with open(converted_path, "rb") as f:
                        docx_data = f.read()
            doc = Document(io.BytesIO(docx_data))
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError("Không đọc được file Word (file có thể bị hỏng hoặc không đúng định dạng .docx).") from e
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("Thiếu thư viện pypdf. Chạy: pip install pypdf")
        try:
            reader = PdfReader(io.BytesIO(data))
        except Exception:
            raise ValueError("Không đọc được file PDF (file có thể bị hỏng hoặc là ảnh scan).")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    raise ValueError("Định dạng không hỗ trợ. Chỉ nhận file .pdf, .doc, .docx, .txt, .md")


def ai_chat(user_prompt, temperature=0.4, max_tokens=4096):
    if not AI_BASE_URL or not AI_API_KEY:
        raise AIError("Chưa cấu hình AI: điền AI_BASE_URL và AI_API_KEY vào file .env rồi khởi động lại server.")
    url = chat_completions_url(AI_BASE_URL)
    payload = {
        "model": AI_MODEL,
        "messages": [
            {"role": "system", "content": "Bạn là trợ lý ra đề trắc nghiệm. Luôn trả về JSON hợp lệ theo đúng yêu cầu."},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": temperature,
        "max_tokens": max_tokens,
    }
    try:
        resp = requests.post(
            url,
            json=payload,
            headers={"Authorization": f"Bearer {AI_API_KEY}", "Content-Type": "application/json"},
            timeout=(15, 300),
        )
    except requests.RequestException as e:
        raise AIError(f"Không kết nối được AI API ({url}): {e}")
    if resp.status_code != 200:
        if "<html" in resp.text[:500].lower() or "<!doctype html" in resp.text[:500].lower():
            raise AIError(f"AI API trả về HTTP {resp.status_code} (dịch vụ trung gian trả trang lỗi). Hãy thử lại sau ít phút.")
        raise AIError(f"AI API trả về HTTP {resp.status_code}: {resp.text[:300]}")
    try:
        return resp.json()["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        raise AIError("Phản hồi của AI không đúng định dạng chat/completions.")


def extract_json_array(content):
    if not content or not content.strip():
        raise AIError("AI trả về nội dung rỗng.")
    text = content.strip()
    fence = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if fence:
        text = fence.group(1)
    start, end = text.find("["), text.rfind("]")
    if start == -1 or end == -1 or end < start:
        raise AIError("Không tìm thấy mảng JSON trong phản hồi của AI.")
    try:
        items = json.loads(text[start:end + 1])
    except ValueError as e:
        raise AIError(f"JSON từ AI không hợp lệ: {e}")
    if not isinstance(items, list):
        raise AIError("Phản hồi của AI không phải là mảng câu hỏi.")
    return items


def normalize_questions(items):
    """Chuẩn hóa + lọc câu hỏi không hợp lệ từ AI."""
    out = []
    for item in items:
        if not isinstance(item, dict):
            continue
        text = str(item.get("text") or item.get("question") or "").strip()
        options = item.get("options") or item.get("choices") or []
        if not isinstance(options, list):
            continue
        options = [str(o).strip() for o in options][:4]
        if not text or len(options) != 4 or any(not o for o in options):
            continue
        raw_answer = str(item.get("answer") or item.get("correct") or "").strip()
        answer = raw_answer.upper()
        if answer in {"A", "B", "C", "D"}:
            correct = options[ord(answer) - ord("A")]
        else:
            # AI có thể trả đáp án dạng văn bản -> khớp cả đúng chữ lẫn không phân biệt hoa thường
            correct = next((o for o in options if o == raw_answer), "")
            if not correct:
                correct = next((o for o in options if o.lower() == raw_answer.lower()), "")
        if not correct:
            continue
        out.append({"text": text, "imageUrl": "", "options": options, "correctText": correct})
    return out


def build_prompt(source, needed, round_no):
    return (
        f"Bạn là giáo viên ra đề. Dựa CHỈ vào tài liệu dưới đây, viết đúng {needed} câu hỏi "
        "trắc nghiệm tiếng Việt, mỗi câu có 4 phương án A/B/C/D và đúng 1 đáp án đúng.\n"
        "Yêu cầu:\n"
        f"- Lần sinh thứ {round_no}: không trùng lặp với câu hỏi đã có (nếu hệ thống yêu cầu thêm).\n"
        "- Phủ toàn bộ nội dung tài liệu, không hỏi thông tin ngoài tài liệu.\n"
        "- Trả về DUY NHẤT một mảng JSON, không giải thích, không markdown, không ký tự ngoài JSON, dạng:\n"
        "[{\"text\": \"Nội dung câu hỏi\", \"options\": [\"A\", \"B\", \"C\", \"D\"], \"answer\": \"A\"}]\n"
        "Trong đó \"answer\" là chữ cái A/B/C/D chỉ phương án đúng.\n\n"
        f"TÀI LIỆU:\n{source}"
    )


def generate_questions(source_text, count=EXAM_QUESTION_COUNT):
    """Gọi AI sinh đúng `count` câu hợp lệ (tự thử lại nếu AI trả thiếu/sai)."""
    source = (source_text or "").strip()[:MAX_SOURCE_CHARS]
    if not source:
        raise AIError("Tài liệu rỗng, không thể sinh đề.")
    collected, seen = [], set()
    last_error = ""
    # Tối đa 8 lô: mỗi lô hỏi 10 câu, cho phép thử lại lô lỗi/thiếu.
    for attempt in range(1, 9):
        # Chia nhỏ để provider không phải xử lý một phản hồi 40 câu quá lớn.
        needed = min(10, count - len(collected))
        if needed <= 0:
            break
        try:
            content = ai_chat(build_prompt(source, needed, attempt))
            raw = extract_json_array(content)
        except AIError as e:
            last_error = str(e)
            continue
        for question in normalize_questions(raw):
            key = question["text"].lower()[:150]
            if key in seen:
                continue
            seen.add(key)
            collected.append(question)
            if len(collected) >= count:
                break
    if len(collected) < count:
        detail = f" ({last_error})" if last_error else ""
        raise AIError(f"AI chỉ sinh được {len(collected)}/{count} câu hợp lệ.{detail}")
    return collected[:count]

@app.get("/", response_class=HTMLResponse)
async def get_home():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()

def _public_quiz(quiz):
    return {
        "id": quiz.get("id"),
        "title": quiz.get("title"),
        "subject": quiz.get("subject"),
        "mode": quiz.get("mode"),
        "question_count": len(quiz.get("questions") or []),
        "created_at": quiz.get("created_at"),
    }


@app.get("/quizzes")
async def list_quiz_sets():
    """Danh sách các bộ đề để học sinh chọn làm (không kèm đáp án)."""
    try:
        quizzes = store.list_quizzes()
    except db.SupabaseError as e:
        return JSONResponse(status_code=503, content={"error": str(e)})
    return {"quizzes": [_public_quiz(q) for q in quizzes]}


@app.get("/quizzes/{quiz_id}/questions")
async def quiz_questions(quiz_id: str):
    quiz = store.get_quiz(quiz_id)
    if not quiz:
        return JSONResponse(status_code=404, content={"error": "Không tìm thấy bộ đề."})
    if quiz.get("mode") == "random":
        return JSONResponse(status_code=503, content={"error": "Bộ đề này sinh ngẫu nhiên theo từng lượt thi."})
    questions = quiz.get("questions") or []
    if not questions:
        return JSONResponse(status_code=503, content={"error": "Bộ đề chưa có câu hỏi."})
    # Never expose answer keys through the public question endpoint.
    return [{k: v for k, v in q.items() if k != "correctText"} for q in questions]

@app.post("/start")
async def start_exam(request: Request):
    try:
        data = await request.json()
        name = (data.get("name") or "").strip()
        quiz_id = (data.get("quiz_id") or "").strip()
        if not name:
            return JSONResponse(status_code=400, content={"error": "Vui lòng nhập họ tên thí sinh!"})
        if len(name) > 50:
            return JSONResponse(status_code=400, content={"error": "Tên không được quá 50 ký tự!"})
        if not quiz_id:
            return JSONResponse(status_code=400, content={"error": "Vui lòng chọn bộ đề trước khi bắt đầu!"})
        quiz = store.get_quiz(quiz_id)
        if not quiz:
            return JSONResponse(status_code=404, content={"error": "Bộ đề không tồn tại hoặc đã bị xóa."})
        # The server owns the answer key and creates the exam order. Never trust
        # question content or correct answers supplied by the browser.
        if quiz.get("mode") == "random":
            try:
                questions = await asyncio.to_thread(generate_questions, quiz.get("source_text") or "")
            except AIError as e:
                return JSONResponse(status_code=502, content={"error": str(e)})
        else:
            questions = quiz.get("questions") or []
            if not questions:
                return JSONResponse(status_code=503, content={"error": "Bộ đề chưa sẵn sàng (chưa có câu hỏi). Vào /admin để sinh đề."})
        # Tạo bản sao để không làm thay đổi kho bộ đề khi xáo trộn
        questions = [{**q, "options": list(q.get("options", []))} for q in questions]
        random.shuffle(questions)
        for question in questions:
            random.shuffle(question["options"])
        cleanup_old_sessions()
        session_id = str(uuid.uuid4())
        start_time = get_current_utc()
        save_session(session_id, name, start_time, questions, quiz_id=quiz_id)
        return {
            "session_id": session_id,
            "start_time": start_time.isoformat(),
            "duration_minutes": EXAM_DURATION_MINUTES,
            "title": quiz.get("title"),
            "subject": quiz.get("subject"),
            "questions": [
                {key: value for key, value in question.items() if key != "correctText"}
                for question in questions
            ]
        }
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})

@app.post("/submit")
async def handle_submit(request: Request):
    try:
        data = await request.json()
        session_id = data.get("session_id")
        if not session_id:
            return JSONResponse(status_code=400, content={"error": "Không tìm thấy phiên làm bài!"})
        session = get_session(session_id)
        if not session:
            return JSONResponse(status_code=400, content={"error": "Phiên làm bài đã hết hạn hoặc không tồn tại!"})
        shuffled_questions = session.get("shuffled_questions") or []
        total = len(shuffled_questions)
        user_answers = data.get("answers", [])
        if not isinstance(user_answers, list) or total == 0 or len(user_answers) != total or any(
            answer is not None and not isinstance(answer, str) for answer in user_answers
        ):
            return JSONResponse(status_code=400, content={"error": "Dữ liệu câu trả lời không hợp lệ!"})
        # Tính điểm dựa trên thứ tự đã xáo trộn
        score = sum(1 for i, q in enumerate(shuffled_questions)
                    if q.get("correctText") and user_answers[i] == q.get("correctText"))
        start_time = datetime.fromisoformat(session["start_time"])
        now = get_current_utc()
        duration_sec = int((now - start_time).total_seconds())
        if duration_sec < 0 or duration_sec > EXAM_DURATION_MINUTES * 60 + 30:
            delete_session(session_id)
            return JSONResponse(status_code=400, content={"error": "Đã hết thời gian làm bài."})
        duration_fmt = format_duration(duration_sec)
        score_display = f"{score}/{total}"
        quiz_title = ""
        if session.get("quiz_id"):
            quiz_title = (store.get_quiz(session["quiz_id"]) or {}).get("title") or ""
        save_submission({
            "id": uuid.uuid4().hex[:12],
            "source": "anonymous",
            "exam_id": None,
            "quiz_id": session.get("quiz_id"),
            "exam_title": quiz_title,
            "user_id": None,
            "user_name": session["name"],
            "score": score,
            "total": total,
            "duration_sec": duration_sec,
            "started_at": session["start_time"],
            "submitted_at": now.isoformat(),
            "answers": user_answers,
        })
        delete_session(session_id)
        return {
            "score": score_display,
            "score_num": score,
            "message": "success",
            "name": session["name"],
            "total_questions": total,
            "duration_formatted": duration_fmt,
            "duration_sec": duration_sec,
            "details": [
                {
                    "idx": i,
                    "text": shuffled_questions[i].get("text", ""),
                    "imageUrl": shuffled_questions[i].get("imageUrl", ""),
                    "options": shuffled_questions[i].get("options", []),
                    "userAnswer": user_answers[i],
                    "correctAnswer": shuffled_questions[i].get("correctText", ""),
                    "isCorrect": (user_answers[i] == shuffled_questions[i].get("correctText", "")) if shuffled_questions[i].get("correctText", "") else False
                }
                for i in range(total)
            ]
        }
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})

# ---------------- Quản trị: tải tài liệu & sinh đề bằng AI ----------------
def admin_authorized(request: Request, admin_key: str) -> bool:
    if not ADMIN_KEY:
        return True
    return admin_key == ADMIN_KEY or request.headers.get("x-admin-key") == ADMIN_KEY


@app.get("/admin")
async def admin_page():
    if not os.path.exists("admin.html"):
        return HTMLResponse("Không tìm thấy admin.html", status_code=404)
    with open("admin.html", "r", encoding="utf-8") as f:
        return HTMLResponse(f.read())


@app.get("/portal")
async def portal_page():
    if not os.path.exists("portal.html"):
        return HTMLResponse("Không tìm thấy portal.html", status_code=404)
    with open("portal.html", "r", encoding="utf-8") as f:
        return HTMLResponse(f.read())


@app.get("/admin/status")
async def admin_status():
    try:
        quizzes = store.list_quizzes()
    except db.SupabaseError as e:
        return JSONResponse(status_code=503, content={"error": str(e)})
    return {
        "ai_configured": bool(AI_BASE_URL and AI_API_KEY),
        "ai_model": AI_MODEL,
        "admin_protected": bool(ADMIN_KEY),
        "quizzes": [{
            **_public_quiz(q),
            "filename": q.get("filename"),
            "source_chars": len(q.get("source_text") or ""),
            "updated_at": q.get("updated_at"),
        } for q in quizzes],
    }


@app.post("/admin/test-api")
async def admin_test_api(request: Request, payload: dict):
    """Chỉ lưu cấu hình sau khi API trả lời thành công."""
    global AI_BASE_URL, AI_API_KEY, AI_MODEL
    if not admin_authorized(request, str(payload.get("admin_key") or "")):
        return JSONResponse(status_code=403, content={"error": "Sai admin key."})
    base_url = str(payload.get("base_url") or "").strip()
    api_key = str(payload.get("api_key") or "").strip()
    model = str(payload.get("model") or AI_MODEL).strip()
    if not base_url or not api_key or not model:
        return JSONResponse(status_code=400, content={"error": "Cần nhập Base URL, API key và model."})
    try:
        response = requests.post(
            chat_completions_url(base_url),
            json={"model": model, "messages": [{"role": "user", "content": "Reply with OK"}], "max_tokens": 8},
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            timeout=(10, 30),
        )
    except requests.RequestException as e:
        return JSONResponse(status_code=502, content={"error": f"Không kết nối được API: {e}"})
    if response.status_code != 200:
        return JSONResponse(status_code=502, content={"error": f"API trả về HTTP {response.status_code}: {response.text[:200]}"})
    try:
        response.json()["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        return JSONResponse(status_code=502, content={"error": "API phản hồi không đúng định dạng chat/completions."})
    AI_BASE_URL, AI_API_KEY, AI_MODEL = base_url, api_key, model
    try:
        save_ai_config(AI_BASE_URL, AI_API_KEY, AI_MODEL)
    except OSError as e:
        return JSONResponse(status_code=500, content={"error": f"Test thành công nhưng không lưu được cấu hình: {e}"})
    return {"message": "API hoạt động và đã lưu cấu hình."}


@app.post("/admin/upload")
async def admin_upload(request: Request, file: UploadFile = File(...),
                       title: str = Form(""), subject: str = Form(""),
                       mode: str = Form("shared"), admin_key: str = Form("")):
    if not admin_authorized(request, admin_key):
        return JSONResponse(status_code=403, content={"error": "Sai admin key."})
    if mode not in ("shared", "random"):
        return JSONResponse(status_code=400, content={"error": "Mode không hợp lệ (chọn shared hoặc random)."})
    data = await file.read()
    if not data:
        return JSONResponse(status_code=400, content={"error": "File rỗng."})
    if len(data) > MAX_UPLOAD_BYTES:
        return JSONResponse(status_code=413, content={"error": "File quá lớn (tối đa 5MB)."})
    try:
        text = extract_text(file.filename or "", data)
    except ValueError as e:
        return JSONResponse(status_code=400, content={"error": str(e)})
    if len(text.strip()) < MIN_SOURCE_CHARS:
        return JSONResponse(status_code=400, content={"error": "Không trích được nội dung đủ dài từ tài liệu (cần ít nhất 200 ký tự)."})

    quiz_title = (title or "").strip() or os.path.splitext(file.filename or "Bộ đề")[0]
    questions = []
    if mode == "shared":
        try:
            questions = await asyncio.to_thread(generate_questions, text)
        except AIError as e:
            # vẫn lưu tài liệu để sinh lại sau
            quiz, _ = store.create_quiz(quiz_title, subject, mode, [], text, file.filename)
            return JSONResponse(status_code=502, content={"error": str(e),
                                "quiz_id": (quiz or {}).get("id")})
    quiz, error = store.create_quiz(quiz_title, subject, mode, questions, text, file.filename)
    if error:
        return JSONResponse(status_code=400, content={"error": error})
    return {
        "message": "success",
        "quiz_id": quiz["id"],
        "title": quiz["title"],
        "subject": quiz["subject"],
        "mode": mode,
        "filename": file.filename,
        "questions": len(questions),
        "source_chars": len(text),
    }


@app.patch("/admin/quizzes/{quiz_id}")
async def admin_update_quiz(request: Request, quiz_id: str, title: str = Form(""),
                            subject: str = Form(""), admin_key: str = Form("")):
    if not admin_authorized(request, admin_key):
        return JSONResponse(status_code=403, content={"error": "Sai admin key."})
    fields = {}
    if title.strip():
        fields["title"] = title.strip()
    if subject.strip():
        fields["subject"] = subject.strip()
    quiz, error = store.update_quiz(quiz_id, **fields)
    if error:
        return JSONResponse(status_code=404, content={"error": error})
    return {"message": "success", "quiz": _public_quiz(quiz)}


@app.delete("/admin/quizzes/{quiz_id}")
async def admin_delete_quiz(request: Request, quiz_id: str, admin_key: str = ""):
    if not admin_authorized(request, admin_key):
        return JSONResponse(status_code=403, content={"error": "Sai admin key."})
    _, error = store.delete_quiz(quiz_id)
    if error:
        return JSONResponse(status_code=404, content={"error": error})
    return {"message": "success"}


@app.get("/leaderboard")
async def get_leaderboard():
    try:
        submissions = store.anonymous_attempts()
    except db.SupabaseError:
        submissions = []
    best = {}
    max_total = 0
    for sub in submissions:
        name = sub.get("user_name") or ""
        if not name:
            continue
        score = int(sub.get("score") or 0)
        dur = int(sub.get("duration_sec") or 0)
        total = int(sub.get("total") or 0)
        max_total = max(max_total, total)
        if (name not in best or score > best[name]["score"]
                or (score == best[name]["score"] and dur < best[name]["duration_sec"])):
            best[name] = {
                "name": name,
                "score": score,
                "duration_sec": dur,
                "duration_formatted": format_duration(dur),
            }
    top3 = sorted(best.values(), key=lambda x: (-x["score"], x["duration_sec"]))[:3]
    return {"leaderboard": top3, "total_questions": max_total or EXAM_QUESTION_COUNT}

@app.get("/history/{name}")
async def get_user_history(name: str):
    try:
        submissions = store.attempts_by_name(name)
    except db.SupabaseError:
        submissions = []
    history = [{
        "name": s.get("user_name"),
        "score": s.get("score"),
        "score_display": f"{s.get('score', 0)}/{s.get('total', 0)}",
        "duration_sec": s.get("duration_sec"),
        "duration_formatted": format_duration(s.get("duration_sec") or 0),
        "submitted_at_display": to_vn_display(s.get("submitted_at")),
        "submitted_at_iso": s.get("submitted_at"),
        "exam_title": s.get("exam_title"),
    } for s in submissions]
    history.sort(key=lambda x: x.get("submitted_at_iso") or "", reverse=True)
    return history

@app.get("/admin-check-history")
async def view_history():
    try:
        return store.anonymous_attempts()
    except db.SupabaseError:
        return []

if __name__ == "__main__":
    import uvicorn
    print("🚀 Đang khởi chạy server Azota Mini tại http://localhost:8000 ...")
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True, ws="none")
